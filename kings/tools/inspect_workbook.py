from pathlib import Path
import json

import openpyxl


downloads = Path.home() / "Downloads"
files = sorted(downloads.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
print(json.dumps([{"name": p.name, "path": str(p), "size": p.stat().st_size} for p in files[:30]], ensure_ascii=True, indent=2))

target = files[0]
for p in files:
    if p.stat().st_size == 13101:
        target = p
        break

print("TARGET", json.dumps({"name": target.name, "path": str(target)}, ensure_ascii=True))
wb = openpyxl.load_workbook(target, data_only=False)
print("SHEETS", json.dumps(wb.sheetnames, ensure_ascii=True))
for ws in wb.worksheets:
    print("SHEET", json.dumps({"title": ws.title, "rows": ws.max_row, "cols": ws.max_column}, ensure_ascii=True))
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), start=1):
        cleaned = [value for value in row]
        if any(value not in (None, "") for value in cleaned):
            print(idx, json.dumps(cleaned, ensure_ascii=False, default=str))
