from __future__ import annotations

from pathlib import Path
import json
import re

import openpyxl


def find_source() -> Path:
    downloads = Path.home() / "Downloads"
    matches = [p for p in downloads.glob("*.xlsx") if "열왕기" in p.name and not p.name.startswith("~$")]
    if matches:
        return max(matches, key=lambda p: p.stat().st_mtime)
    sized = [p for p in downloads.glob("*.xlsx") if p.stat().st_size == 13101 and not p.name.startswith("~$")]
    if sized:
        return max(sized, key=lambda p: p.stat().st_mtime)
    raise FileNotFoundError("열왕기_재위기간.xlsx 파일을 찾지 못했습니다.")


def parse_duration(name: str) -> float | None:
    match = re.search(r"\(([\d.]+)\)", name or "")
    return float(match.group(1)) if match else None


def clean_name(name: str) -> str:
    return re.sub(r"\s*\([^)]*\)\s*$", "", name or "").strip()


def fmt_years(value: float | None) -> str:
    if value is None:
        return "기록 없음"
    if value == 0:
        return "수개월"
    if value < 1:
        months = round(value * 12)
        return f"약 {months}개월"
    if value == int(value):
        return f"{int(value)}년"
    return f"{value:g}년"


def build_data(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rulers = []
    events = []
    for row_idx in range(1, ws.max_row + 1):
        judah = ws.cell(row_idx, 1).value
        israel = ws.cell(row_idx, 5).value
        note = ws.cell(row_idx, 7).value
        if isinstance(judah, str) and row_idx > 1:
            duration = parse_duration(judah)
            rulers.append(
                {
                    "kingdom": "유다",
                    "name": clean_name(judah),
                    "label": judah,
                    "duration": duration,
                    "durationText": fmt_years(duration),
                    "row": row_idx,
                }
            )
        if isinstance(israel, str) and row_idx > 1 and israel.strip():
            duration = parse_duration(israel)
            rulers.append(
                {
                    "kingdom": "이스라엘",
                    "name": clean_name(israel),
                    "label": israel,
                    "duration": duration,
                    "durationText": fmt_years(duration),
                    "row": row_idx,
                }
            )
        if isinstance(note, str) and note.strip():
            events.append({"row": row_idx, "note": note.strip()})

    rows = []
    for row_idx in range(1, ws.max_row + 1):
        rows.append(
            {
                "row": row_idx,
                "cells": [
                    ws.cell(row_idx, col_idx).value
                    for col_idx in range(1, ws.max_column + 1)
                ],
            }
        )

    def build_cards(kingdom: str) -> list[dict]:
        if kingdom == "유다":
            name_col = 1
            year_col = 2
        else:
            name_col = 5
            year_col = 4

        starts = []
        for row_idx in range(2, ws.max_row + 1):
            name = ws.cell(row_idx, name_col).value
            if isinstance(name, str) and name.strip():
                starts.append(row_idx)

        cards = []
        for index, start in enumerate(starts):
            next_start = starts[index + 1] if index + 1 < len(starts) else ws.max_row + 1
            end = next_start - 1
            name = ws.cell(start, name_col).value
            year_values = []
            years = []
            for row_idx in range(start, end + 1):
                year = ws.cell(row_idx, year_col).value
                if year not in (None, ""):
                    year_values.append(year)
                    years.append({"row": row_idx, "value": year})
            start_year = year_values[0] if year_values else None
            end_year = year_values[-1] if year_values else None
            duration = parse_duration(name)
            if duration is not None and duration != int(duration):
                if years:
                    years[-1]["value"] = duration
                    year_values[-1] = duration
                    end_year = duration
                else:
                    years.append({"row": start, "value": duration})
                    start_year = duration
                    end_year = duration
            cards.append(
                {
                    "kingdom": kingdom,
                    "name": name,
                    "startRow": start,
                    "endRow": end,
                    "startYear": start_year,
                    "endYear": end_year,
                    "years": years,
                    "duration": duration,
                    "tone": index % 2,
                }
            )
        return cards

    cards = {
        "judah": build_cards("유다"),
        "israel": build_cards("이스라엘"),
    }

    for ruler in rulers:
        row_notes = [event["note"] for event in events if event["row"] == ruler["row"]]
        ruler["notes"] = row_notes

    by_kingdom = {
        "유다": [r for r in rulers if r["kingdom"] == "유다"],
        "이스라엘": [r for r in rulers if r["kingdom"] == "이스라엘"],
    }
    durations = [r["duration"] for r in rulers if isinstance(r["duration"], (int, float))]
    longest = max(rulers, key=lambda r: r["duration"] or 0)
    return {
        "source": path.name,
        "sheet": ws.title,
        "maxRow": ws.max_row,
        "rulers": rulers,
        "events": events,
        "rows": rows,
        "cards": cards,
        "summary": {
            "total": len(rulers),
            "judah": len(by_kingdom["유다"]),
            "israel": len(by_kingdom["이스라엘"]),
            "average": round(sum(durations) / len(durations), 1),
            "longest": {
                "name": longest["name"],
                "kingdom": longest["kingdom"],
                "durationText": longest["durationText"],
            },
        },
    }


def html_page(data: dict) -> str:
    data_json = json.dumps(data, ensure_ascii=False)
    return f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>열왕기 재위기간</title>
  <style>
    :root {{
      --ink: #201a17;
      --muted: #6d635d;
      --paper: #f8f4ec;
      --paper-2: #fffaf1;
      --line: rgba(74, 58, 44, .18);
      --judah: #9b3f2f;
      --judah-soft: #f3d6ca;
      --israel: #1f6f78;
      --israel-soft: #cfe8e7;
      --gold: #b88924;
      --shadow: 0 18px 50px rgba(55, 41, 30, .13);
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      color: var(--ink);
      background:
        radial-gradient(circle at 16% 10%, rgba(184, 137, 36, .14), transparent 28rem),
        linear-gradient(135deg, #f4efe5 0%, #fffaf2 52%, #edf5f3 100%);
      font-family: "Pretendard Variable", Pretendard, "SUIT Variable", SUIT, "Noto Sans KR", "Segoe UI", "Apple SD Gothic Neo", "Malgun Gothic", system-ui, sans-serif;
      line-height: 1.5;
    }}
    .shell {{ max-width: 1180px; margin: 0 auto; padding: 22px 16px 44px; }}
    header {{
      padding: 6px 0 14px;
    }}
    h1 {{
      margin: 0;
      font-size: clamp(2.1rem, 5vw, 4.6rem);
      line-height: .98;
      letter-spacing: 0;
      font-weight: 850;
    }}
    .board-wrap {{
      margin-top: 8px;
      overflow: auto;
      max-height: 84vh;
      border: 1px solid rgba(28, 24, 20, .22);
      background: #f7f5ea;
      box-shadow: var(--shadow);
    }}
    .board {{
      min-width: 860px;
      position: relative;
      display: grid;
      grid-template-columns: 1fr 1fr;
      column-gap: 18px;
      padding: 2px;
    }}
    .board-head {{
      grid-column: 1 / -1;
      display: grid;
      grid-template-columns: 1fr 1fr;
      column-gap: 18px;
      align-items: center;
      height: 34px;
      margin-bottom: 2px;
      color: #5f554e;
      font-size: .9rem;
      font-weight: 650;
      text-align: center;
      border-bottom: 1px solid rgba(28, 24, 20, .42);
    }}
    .lane {{
      position: relative;
      min-height: calc(var(--rows, 120) * var(--row-h, 32px));
    }}
    .lane.judah {{ grid-column: 1; }}
    .lane.israel {{ grid-column: 2; }}
    .king-card {{
      position: absolute;
      left: 0;
      right: 0;
      min-height: 22px;
      padding: 7px 10px 12px;
      border: 1px solid #1c1814;
      box-shadow: none;
      transition: opacity .14s ease, transform .14s ease, filter .14s ease;
      overflow: hidden;
    }}
    .king-card:hover {{ transform: translateY(-1px); filter: saturate(1.08); }}
    .king-card.dim {{ opacity: .16; }}
    .king-card.tone-0.judah {{
      color: #241915;
      background: linear-gradient(135deg, #f7dcc9 0%, #f0c1a6 100%);
      border-color: #8f513a;
    }}
    .king-card.tone-1.judah {{
      color: #221b10;
      background: linear-gradient(135deg, #ffe8a8 0%, #f3c862 100%);
      border-color: #9b7421;
    }}
    .king-card.tone-0.israel {{
      color: #102524;
      background: linear-gradient(135deg, #cfece7 0%, #8ed0c8 100%);
      border-color: #327a75;
    }}
    .king-card.tone-1.israel {{
      color: #152213;
      background: linear-gradient(135deg, #d9ecd0 0%, #a7d38e 100%);
      border-color: #5d843e;
    }}
    .king-card.israel {{
      text-align: right;
    }}
    .king-name {{
      position: absolute;
      top: 6px;
      left: 10px;
      right: 64px;
      width: fit-content;
      max-width: calc(100% - 66px);
      padding: 2px 6px;
      background: rgba(80, 48, 28, .12);
      border-radius: 3px;
      font-weight: 650;
      overflow-wrap: anywhere;
      line-height: 1.2;
    }}
    .king-card.israel .king-name {{
      left: 64px;
      right: 10px;
      margin-left: auto;
      background: rgba(20, 76, 73, .12);
    }}
    .year-marker {{
      position: absolute;
      right: 10px;
      width: 38px;
      text-align: center;
      color: #000;
      font-family: "Segoe UI", "Pretendard Variable", Pretendard, "Noto Sans KR", sans-serif;
      font-weight: 600;
      font-variant-numeric: tabular-nums;
      line-height: 1.1;
      z-index: 2;
    }}
    .king-card.israel .year-marker {{
      left: 18px;
      right: auto;
      text-align: center;
    }}
    .year-arrow {{
      position: absolute;
      right: 28.5px;
      width: 1px;
      background: rgba(20, 17, 14, .48);
      z-index: 1;
    }}
    .king-card.israel .year-arrow {{
      left: 36.5px;
      right: auto;
    }}
    .year-arrow::after {{
      content: "";
      position: absolute;
      left: -4px;
      bottom: -1px;
      width: 0;
      height: 0;
      border-left: 4px solid transparent;
      border-right: 4px solid transparent;
      border-top: 7px solid rgba(20, 17, 14, .62);
    }}
    @media (max-width: 860px) {{
      .shell {{ padding-inline: 14px; }}
      .board {{ min-width: 680px; column-gap: 12px; padding-inline: 2px; }}
      .board-head {{ column-gap: 12px; }}
    }}
  </style>
</head>
<body>
  <main class="shell">
    <header>
      <h1>열왕기 재위기간</h1>
    </header>

    <section class="board-wrap" aria-label="왕국별 재위기간 비교 보드">
      <div class="board" id="board">
        <div class="board-head">
          <strong>유다(예루살렘)</strong>
          <strong>이스라엘(사마리아)</strong>
        </div>
        <div class="lane judah" id="lane-judah"></div>
        <div class="lane israel" id="lane-israel"></div>
      </div>
    </section>
  </main>

  <script>
    const DATA = {data_json};
    const board = document.querySelector("#board");
    const laneJudah = document.querySelector("#lane-judah");
    const laneIsrael = document.querySelector("#lane-israel");
    const rowHeight = 32;
    const blankRowHeight = 3;
    const rowWeights = new Map(DATA.rows.map(row => [
      row.row,
      row.cells.some(value => value !== null && value !== "") ? rowHeight : blankRowHeight
    ]));

    function visualY(rowNumber) {{
      let y = 0;
      for (let row = 2; row < rowNumber; row += 1) {{
        y += rowWeights.get(row) ?? rowHeight;
      }}
      return y;
    }}

    function displayValue(value) {{
      if (value == null) return "";
      if (typeof value === "number" && Number.isInteger(value)) return String(value);
      return String(value);
    }}

    function escapeHtml(value) {{
      return String(value).replace(/[&<>"']/g, ch => ({{ "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#39;" }}[ch]));
    }}

    function renderCard(card) {{
      const top = Math.max(0, visualY(card.startRow));
      const yearRows = (card.years || []).map(item => item.row);
      const visualEndRow = yearRows.length ? Math.max(...yearRows) : card.startRow;
      const height = Math.max(32, visualY(visualEndRow) - visualY(card.startRow) + 30);
      const kingdomClass = card.kingdom === "유다" ? "judah" : "israel";
      const markers = (card.years || []).map(item => {{
        const markerTop = Math.max(8, visualY(item.row) - visualY(card.startRow) + 8);
        return `<span class="year-marker" style="top:${{markerTop}}px">${{escapeHtml(displayValue(item.value))}}</span>`;
      }}).join("");
      const markerPositions = (card.years || []).map(item => Math.max(8, visualY(item.row) - visualY(card.startRow) + 8));
      const arrows = markerPositions.slice(0, -1).map((pos, index) => {{
        const next = markerPositions[index + 1];
        return `<span class="year-arrow" style="top:${{pos + 19}}px; height:${{Math.max(8, next - pos - 24)}}px"></span>`;
      }}).join("");
      return `
        <article class="king-card ${{kingdomClass}} tone-${{card.tone}}" data-kingdom="${{escapeHtml(card.kingdom)}}" data-name="${{escapeHtml(card.name)}}" style="top:${{top}}px; height:${{height}}px">
          <div class="king-name">${{escapeHtml(card.name)}}</div>
          ${{arrows}}
          ${{markers}}
        </article>
      `;
    }}

    function renderBoard() {{
      const rowCount = Math.max(...DATA.rows.map(row => row.row)) - 1;
      board.style.setProperty("--rows", rowCount);
      board.style.setProperty("--row-h", `${{rowHeight}}px`);
      board.style.minHeight = `${{visualY(rowCount + 2)}}px`;
      laneJudah.innerHTML = DATA.cards.judah.map(renderCard).join("");
      laneIsrael.innerHTML = DATA.cards.israel.map(renderCard).join("");
    }}
    renderBoard();
  </script>
</body>
</html>
"""


def main() -> None:
    source = find_source()
    data = build_data(source)
    output_dir = Path.cwd() / "outputs"
    output_dir.mkdir(exist_ok=True)
    output = output_dir / "kings_reign_timeline.html"
    output.write_text(html_page(data), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
